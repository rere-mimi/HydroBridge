"""Single HydroBridge results workbook with native Excel charts."""

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import AreaChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="0F172A")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="0F172A")
LABEL_FONT = Font(name="Calibri", size=11, bold=True, color="334155")
BODY_FONT = Font(name="Calibri", size=11, color="0F172A")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="0F172A")
ALT_FILL = PatternFill("solid", fgColor="F8FAFC")
BLUE_FILL = PatternFill("solid", fgColor="DBEAFE")
THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
WRAP = Alignment(wrap_text=True, vertical="top")

METHODOLOGY = (
    "HydroBridge raises a water surface on each transect until Manning discharge "
    "matches the design flow for that return period:\n\n"
    "    Q = (1/n) · A · R^(2/3) · S^(1/2)\n\n"
    "A (flow area), P (wetted perimeter), and R = A/P (hydraulic radius) are integrated "
    "from trapezoidal segments between DEM samples. Only the wet interval that is "
    "continuously connected to the main channel (the centreline crossing at the "
    "transect midpoint) is included. Isolated depressions below the water surface "
    "that are separated by high ground do not contribute to A, P, R, top width, "
    "depth, velocity, or discharge until the intervening ridge is overtopped.\n\n"
    "Assumptions: steady uniform flow on each transect; a single composite section "
    "(no left/right overbank subdivision); bed slope S is |dz/ds| from a linear fit "
    "to centreline DEM samples and is shared by every transect; Manning’s n is constant; "
    "the drawn centreline forward direction is downstream. This is a screening tool, "
    "not a replacement for a calibrated 1D/2D model."
)


def _hex(color, fallback="0284C7"):
    text = str(color or fallback).replace("#", "").upper()
    return text if len(text) == 6 else fallback


def _num(value):
    if value is None:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number != number or number in (float("inf"), float("-inf")):
        return None
    return number


def _text(value, default=""):
    if value is None:
        return default
    return str(value)


def _sheet_name(transect):
    name = f"Transect {int(transect)}"
    return name[:31]


def _write_kv(ws, row, label, value, *, value_col=2):
    cell = ws.cell(row, 1, label)
    cell.font = LABEL_FONT
    cell.alignment = Alignment(vertical="center")
    data = ws.cell(row, value_col, value)
    data.font = BODY_FONT
    data.alignment = Alignment(vertical="center", wrap_text=True)
    return row + 1


def _header_row(ws, row, headers, start_col=1):
    for offset, title in enumerate(headers):
        cell = ws.cell(row, start_col + offset, title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        cell.border = THIN


def _style_table_cell(cell, alt=False):
    cell.font = BODY_FONT
    cell.border = THIN
    cell.alignment = Alignment(vertical="center")
    if alt:
        cell.fill = ALT_FILL


def _status_text(hyd):
    if not hyd:
        return "—"
    if hyd.get("overtopped"):
        return "Overtops banks"
    if hyd.get("conveys"):
        return "OK"
    return "Cannot convey"


def _dem_source_label(layout):
    source = (layout or {}).get("dem_source")
    if source == "linz-lidar-1m":
        tiles = layout.get("tiles") or []
        sheet = f" (tiles {', '.join(tiles)})" if tiles else ""
        return f"New Zealand LiDAR 1 m DEM, LINZ layer 121859{sheet}"
    if source == "wcs":
        return "User-supplied WCS raster"
    if source == "local":
        name = (layout or {}).get("dem_file") or "Local GeoTIFF"
        return f"Local DEM ({name})"
    return source or "Not recorded"


def _centerline_label(layout, project):
    source = (project or {}).get("centerline_source") or (layout or {}).get("centerline_source")
    labels = {
        "drawn": "User-drawn river centreline (50 m upstream and downstream buffers along the channel)",
        "osm": "OpenStreetMap waterway near the bridge pin",
        "synthetic": "Synthetic east–west line through the pin (no OSM waterway found)",
    }
    return labels.get(source, source or "River centreline")


def _seed_index(distances):
    finite = [(i, float(x)) for i, x in enumerate(distances) if _num(x) is not None]
    if not finite:
        return 0
    mid = 0.5 * (finite[0][1] + finite[-1][1])
    return min(finite, key=lambda item: abs(item[1] - mid))[0]


def _interp_y(xs, ys, x):
    points = [(_num(a), _num(b)) for a, b in zip(xs, ys)]
    points = [(a, b) for a, b in points if a is not None and b is not None]
    if not points:
        return None
    if x is None:
        return None
    points.sort()
    if x <= points[0][0]:
        return points[0][1]
    if x >= points[-1][0]:
        return points[-1][1]
    for (x0, y0), (x1, y1) in zip(points, points[1:]):
        if x0 <= x <= x1:
            if x1 == x0:
                return y0
            t = (x - x0) / (x1 - x0)
            return y0 + t * (y1 - y0)
    return points[-1][1]


def _style_fill_series(series, color, *, filled=True):
    gp = series.graphicalProperties
    if filled:
        gp.solidFill = color
        gp.line.solidFill = color
    else:
        gp.noFill = True
        gp.line.noFill = True


def _style_line_series(series, color, *, width=18000, dash=None, marker=None):
    series.graphicalProperties.line.solidFill = color
    series.graphicalProperties.line.width = width
    if dash:
        series.graphicalProperties.line.dashStyle = dash
    if marker:
        series.marker = marker
    else:
        series.marker = Marker(symbol=None)


def write_screening_workbook(
    path,
    summary_rows=None,
    data_rows=None,
    *,
    project=None,
    layout=None,
    scenarios=None,
    transects=None,
    centerline_profile=None,
):
    """Write one results workbook: Summary + one sheet per transect, native charts only."""
    path = Path(path)
    project = dict(project or {})
    layout = dict(layout or {})
    scenarios = list(scenarios or [])
    transects = list(transects or [])
    if not transects and summary_rows:
        transects = list(summary_rows)
    centerline_profile = dict(centerline_profile or {})

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    _write_summary_sheet(summary, project, layout, scenarios, transects, centerline_profile)
    for item in transects:
        name = _sheet_name(item.get("transect") or 1)
        n = 2
        base = name
        while name in wb.sheetnames:
            name = f"{base[:27]} {n}"[:31]
            n += 1
        ws = wb.create_sheet(name)
        _write_transect_sheet(ws, item, project, layout, scenarios)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _write_summary_sheet(ws, project, layout, scenarios, transects, centerline_profile):
    ws["A1"] = "HydroBridge results"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")
    ws["A2"] = "One workbook for the screening run. Charts are native Excel charts linked to the tables on this sheet and each transect sheet. Edit a water level or elevation and the plots update."
    ws["A2"].font = BODY_FONT
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:F3")
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 22

    row = 5
    ws.cell(row, 1, "Project information").font = SECTION_FONT
    row += 1
    created = project.get("created") or datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    row = _write_kv(ws, row, "Bridge", project.get("bridge_name") or "Unnamed bridge")
    row = _write_kv(ws, row, "Latitude", _num(project.get("lat")))
    row = _write_kv(ws, row, "Longitude", _num(project.get("lon")))
    row = _write_kv(ws, row, "Run time", created)
    row = _write_kv(ws, row, "Analysis length (m)", _num(layout.get("along_m")))
    row = _write_kv(ws, row, "Transects", _num(layout.get("n_transects")) or len(transects))
    row = _write_kv(ws, row, "Transect spacing (m)", _num(layout.get("interval_m")))
    row = _write_kv(ws, row, "Transect length (m)", _num(layout.get("transect_length_m")))
    row = _write_kv(ws, row, "Sample spacing (m)", _num(layout.get("sample_spacing_m")))

    row += 1
    ws.cell(row, 1, "Data sources").font = SECTION_FONT
    row += 1
    row = _write_kv(ws, row, "Elevation (DEM)", _dem_source_label(layout))
    row = _write_kv(
        ws,
        row,
        "DEM clip window (m)",
        (
            f"{layout.get('upstream_m')} upstream / {layout.get('downstream_m')} downstream / "
            f"{layout.get('lateral_m')} each side"
            if layout.get("upstream_m") is not None
            else "—"
        ),
    )
    row = _write_kv(ws, row, "River centreline", _centerline_label(layout, project))
    row = _write_kv(ws, row, "Survey data", "None — elevations sampled from the DEM along generated transects")
    row = _write_kv(ws, row, "Roughness", f"Manning’s n = {layout.get('mannings_n')}")
    slope = _num(layout.get("slope"))
    row = _write_kv(
        ws,
        row,
        "Energy slope S",
        f"{slope:.6f} (absolute centreline bed slope from DEM)" if slope is not None else "—",
    )
    if scenarios:
        hydrology = "; ".join(
            f"{item.get('label')}: Q = {item.get('flow_m3_s')} m³/s" for item in scenarios
        )
    else:
        hydrology = f"Q = {layout.get('flow_m3_s')} m³/s"
    row = _write_kv(ws, row, "Hydrological inputs", hydrology)

    row += 1
    ws.cell(row, 1, "Hydraulic methodology").font = SECTION_FONT
    row += 1
    start = row
    ws.cell(row, 1, METHODOLOGY)
    ws.cell(row, 1).font = BODY_FONT
    ws.cell(row, 1).alignment = WRAP
    ws.merge_cells(start_row=row, start_column=1, end_row=row + 7, end_column=6)
    for r in range(row, row + 8):
        ws.row_dimensions[r].height = 18
    row += 9

    ws.cell(row, 1, "Hydraulic results").font = SECTION_FONT
    row += 1
    result_headers = [
        "Transect",
        "Offset_m",
        "Station_m",
        "Return period",
        "Years",
        "Flow_m3s",
        "Water_level_m",
        "Max_depth_m",
        "Width_m",
        "Area_m2",
        "Wetted_perimeter_m",
        "Hydraulic_radius_m",
        "Velocity_ms",
        "Discharge_m3s",
        "Status",
        "Sheet",
    ]
    _header_row(ws, row, result_headers)
    ws.row_dimensions[row].height = 28
    row += 1
    table_start = row
    alt = False
    for item in transects:
        aris = item.get("aris") or {}
        rows_for = scenarios or [{"key": key, "label": key, "years": key, "flow_m3_s": None} for key in aris]
        if not rows_for:
            rows_for = [{}]
        for scenario in rows_for:
            hyd = aris.get(scenario.get("key")) or item
            values = [
                item.get("transect"),
                _num(item.get("offset_m")),
                _num(item.get("station_m")),
                scenario.get("label") or "ARI",
                scenario.get("years"),
                _num(scenario.get("flow_m3_s") if scenario.get("flow_m3_s") is not None else hyd.get("flow_m3_s")),
                _num(hyd.get("water_level_m")),
                _num(hyd.get("max_depth_m")),
                _num(hyd.get("width_m")),
                _num(hyd.get("area_m2")),
                _num(hyd.get("wetted_perimeter_m")),
                _num(hyd.get("hydraulic_radius_m")),
                _num(hyd.get("velocity_m_s")),
                _num(hyd.get("discharge_m3_s")),
                _status_text(hyd),
                _sheet_name(item.get("transect") or 1),
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row, col, value)
                _style_table_cell(cell, alt)
            row += 1
            alt = not alt
    table_end = row - 1

    row += 2
    ws.cell(row, 1, "Longitudinal profile data").font = SECTION_FONT
    row += 1
    ws.cell(
        row,
        1,
        "River-bed samples along the centreline. Transect water levels sit at the cross-section stations. Charts read these columns.",
    ).font = BODY_FONT
    row += 1
    bed_headers = ["Station_m", "Bed_elevation_m"]
    _header_row(ws, row, bed_headers)
    bed_header_row = row
    row += 1
    bed_start = row
    dists = centerline_profile.get("distance_m") or []
    elevs = centerline_profile.get("elevation_m") or []
    for dist, elev in zip(dists, elevs):
        station = _num(dist)
        if station is not None:
            station = station  # already along the reach from its start
        ws.cell(row, 1, station)
        ws.cell(row, 2, _num(elev))
        _style_table_cell(ws.cell(row, 1))
        _style_table_cell(ws.cell(row, 2))
        row += 1
    bed_end = row - 1
    if bed_end < bed_start:
        ws.cell(bed_start, 1, 0)
        ws.cell(bed_start, 2, None)
        bed_end = bed_start

    xs_col = 4
    xs_headers = ["Transect", "Station_m", "Bed_elevation_m"]
    for scenario in scenarios:
        xs_headers.append(f"{scenario.get('label')} WL_m")
    _header_row(ws, bed_header_row, xs_headers, start_col=xs_col)
    xs_start = bed_start
    xs_row = xs_start
    for item in transects:
        station = _num(item.get("station_m"))
        bed = _interp_y(dists, elevs, station)
        ws.cell(xs_row, xs_col, item.get("transect"))
        ws.cell(xs_row, xs_col + 1, station)
        ws.cell(xs_row, xs_col + 2, bed)
        aris = item.get("aris") or {}
        for offset, scenario in enumerate(scenarios):
            hyd = aris.get(scenario.get("key")) or {}
            ws.cell(xs_row, xs_col + 3 + offset, _num(hyd.get("water_level_m")))
        for col in range(xs_col, xs_col + 3 + max(len(scenarios), 0)):
            _style_table_cell(ws.cell(xs_row, col))
        xs_row += 1
    xs_end = max(xs_row - 1, xs_start)

    chart = ScatterChart()
    chart.title = "Longitudinal profile"
    chart.style = 10
    chart.x_axis.title = "Station along the river (m)"
    chart.y_axis.title = "Elevation (m)"
    chart.height = 10
    chart.width = 18
    chart.legend.position = "b"
    x_bed = Reference(ws, min_col=1, min_row=bed_start, max_row=bed_end)
    y_bed = Reference(ws, min_col=2, min_row=bed_start, max_row=bed_end)
    bed_series = Series(y_bed, x_bed, title="River bed")
    _style_line_series(bed_series, "0F172A", width=20000)
    chart.series.append(bed_series)

    x_xs = Reference(ws, min_col=xs_col + 1, min_row=xs_start, max_row=xs_end)
    y_xs = Reference(ws, min_col=xs_col + 2, min_row=xs_start, max_row=xs_end)
    xs_series = Series(y_xs, x_xs, title="Cross-section location")
    xs_series.marker = Marker(symbol="diamond", size=8)
    xs_series.marker.graphicalProperties.solidFill = "EA580C"
    xs_series.marker.graphicalProperties.line.solidFill = "9A3412"
    xs_series.graphicalProperties.line.noFill = True
    chart.series.append(xs_series)

    for offset, scenario in enumerate(scenarios):
        y_wl = Reference(ws, min_col=xs_col + 3 + offset, min_row=xs_start, max_row=xs_end)
        series = Series(y_wl, x_xs, title=f"{scenario.get('label')} water surface")
        _style_line_series(series, _hex(scenario.get("color")), width=16000, dash="dash")
        series.marker = Marker(symbol="circle", size=6)
        series.marker.graphicalProperties.solidFill = _hex(scenario.get("color"))
        chart.series.append(series)

    ws.add_chart(chart, "H5")

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18
    for col in range(3, 17):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.freeze_panes = f"A{table_start}"
    ws.print_title_rows = f"1:{table_start - 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddHeader.left.text = "HydroBridge screening"
    return table_end


def _write_transect_sheet(ws, item, project, layout, scenarios):
    transect = int(item.get("transect") or 1)
    coords = item.get("coords") or []
    start = coords[0] if coords else [None, None]
    end = coords[-1] if len(coords) > 1 else start
    distances = list(item.get("distance_m") or [])
    elevations = list(item.get("elevation_m") or [])
    lons = list(item.get("longitude") or [])
    lats = list(item.get("latitude") or [])
    n = max(len(distances), 1)
    while len(elevations) < n:
        elevations.append(None)
    while len(lons) < n:
        lons.append(None)
    while len(lats) < n:
        lats.append(None)

    ws["A1"] = f"Transect {transect}"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")
    ws["A2"] = (
        "Edit Water_level_m in the results table to restage the chart. "
        "Blue fill is the connected flow area only; disconnected depressions stay unfilled."
    )
    ws["A2"].font = BODY_FONT
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:F3")

    row = 5
    ws.cell(row, 1, "Transect metadata").font = SECTION_FONT
    row += 1
    row = _write_kv(ws, row, "Bridge", project.get("bridge_name") or "Unnamed bridge")
    row = _write_kv(ws, row, "Offset from pin (m)", _num(item.get("offset_m")))
    row = _write_kv(ws, row, "Station along river (m)", _num(item.get("station_m")))
    row = _write_kv(ws, row, "Start longitude", _num(start[0]) if start else None)
    row = _write_kv(ws, row, "Start latitude", _num(start[1]) if start else None)
    row = _write_kv(ws, row, "End longitude", _num(end[0]) if end else None)
    row = _write_kv(ws, row, "End latitude", _num(end[1]) if end else None)
    row = _write_kv(ws, row, "Samples", item.get("n_samples") or n)
    row = _write_kv(ws, row, "Sample spacing (m)", _num(item.get("sample_spacing_m") or layout.get("sample_spacing_m")))
    row = _write_kv(ws, row, "Manning’s n", _num(layout.get("mannings_n")))
    row = _write_kv(ws, row, "Bed slope S", _num(layout.get("slope")))
    seed_cell = ws.cell(row, 1, "Channel seed row")
    seed_cell.font = LABEL_FONT
    seed_value_cell = ws.cell(row, 2)
    seed_value_cell.font = BODY_FONT
    seed_row_cell = f"${seed_value_cell.column_letter}${seed_value_cell.row}"
    row += 2

    ws.cell(row, 1, "Hydraulic results by return period").font = SECTION_FONT
    row += 1
    hyd_headers = [
        "Return period",
        "Years",
        "Flow_m3s",
        "Water_level_m",
        "Max_depth_m",
        "Width_m",
        "Area_m2",
        "Wetted_perimeter_m",
        "Hydraulic_radius_m",
        "Velocity_ms",
        "Discharge_m3s",
        "Status",
    ]
    _header_row(ws, row, hyd_headers)
    hyd_header_row = row
    row += 1
    hyd_start = row
    aris = item.get("aris") or {}
    rows_for = scenarios or [{"key": key, "label": key, "years": key, "flow_m3_s": None, "color": "#0284C7"} for key in aris]
    if not rows_for:
        rows_for = [{"key": "100y", "label": "100-year ARI", "years": 100, "flow_m3_s": None, "color": "#7c3aed"}]
    for scenario in rows_for:
        hyd = aris.get(scenario.get("key")) or {}
        values = [
            scenario.get("label") or "ARI",
            scenario.get("years"),
            _num(scenario.get("flow_m3_s") if scenario.get("flow_m3_s") is not None else hyd.get("flow_m3_s")),
            _num(hyd.get("water_level_m")),
            _num(hyd.get("max_depth_m")),
            _num(hyd.get("width_m")),
            _num(hyd.get("area_m2")),
            _num(hyd.get("wetted_perimeter_m")),
            _num(hyd.get("hydraulic_radius_m")),
            _num(hyd.get("velocity_m_s")),
            _num(hyd.get("discharge_m3_s")),
            _status_text(hyd),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            _style_table_cell(cell)
            if col == 4:
                cell.fill = BLUE_FILL
        row += 1
    hyd_end = row - 1
    primary_wl = f"$D${hyd_start}"

    row += 2
    ws.cell(row, 1, "Cross-section geometry").font = SECTION_FONT
    row += 1
    geom_headers = [
        "Distance_m",
        "Elevation_m",
        "Longitude",
        "Latitude",
        "Connected",
        "Fill_base_m",
        "Connected_flow_area_m",
    ]
    for scenario in rows_for:
        geom_headers.append(f"{scenario.get('label')} WL_m")
    _header_row(ws, row, geom_headers)
    geom_header_row = row
    row += 1
    geom_start = row
    seed_idx = _seed_index(distances)
    ws[seed_row_cell] = geom_start + seed_idx
    for i in range(n):
        excel_row = geom_start + i
        ws.cell(excel_row, 1, _num(distances[i]))
        ws.cell(excel_row, 2, _num(elevations[i]))
        ws.cell(excel_row, 3, _num(lons[i]) if i < len(lons) else None)
        ws.cell(excel_row, 4, _num(lats[i]) if i < len(lats) else None)
        connected = (
            f'IF(OR(NOT(ISNUMBER($B{excel_row})),$B{excel_row}>={primary_wl}),FALSE,'
            f'MAX(INDEX($B${geom_start}:$B${geom_start + n - 1},'
            f'MIN(ROW()-{geom_start - 1},{seed_row_cell}-{geom_start - 1})):'
            f'INDEX($B${geom_start}:$B${geom_start + n - 1},'
            f'MAX(ROW()-{geom_start - 1},{seed_row_cell}-{geom_start - 1})))<{primary_wl})'
        )
        ws.cell(excel_row, 5, f"={connected}")
        ws.cell(excel_row, 6, f'=IF(ISNUMBER($B{excel_row}),$B{excel_row},0)')
        ws.cell(excel_row, 7, f'=IF(E{excel_row},MAX({primary_wl}-$B{excel_row},0),0)')
        for offset, _scenario in enumerate(rows_for):
            wl_cell = f"$D${hyd_start + offset}"
            ws.cell(excel_row, 8 + offset, f'=IF(ISNUMBER($B{excel_row}),{wl_cell},NA())')
        for col in range(1, 8 + len(rows_for)):
            _style_table_cell(ws.cell(excel_row, col))
    geom_end = geom_start + n - 1

    area = AreaChart()
    area.grouping = "stacked"
    area.title = f"Transect {transect} cross-section"
    area.style = 10
    area.y_axis.title = "Elevation (m)"
    area.x_axis.title = "Distance across the river (m)"
    area.height = 10
    area.width = 18
    area.legend.position = "b"
    cats = Reference(ws, min_col=1, min_row=geom_start, max_row=geom_end)
    fill_data = Reference(ws, min_col=6, min_row=geom_header_row, max_col=7, max_row=geom_end)
    area.add_data(fill_data, titles_from_data=True)
    area.set_categories(cats)
    if len(area.series) >= 2:
        area.series[0].tx = SeriesLabel(v="")
        _style_fill_series(area.series[0], "F8FAFC", filled=False)
        area.series[1].tx = SeriesLabel(v="Connected flow area")
        _style_fill_series(area.series[1], "2563EB", filled=True)

    line = LineChart()
    ground = Reference(ws, min_col=2, min_row=geom_header_row, max_row=geom_end)
    line.add_data(ground, titles_from_data=True)
    if rows_for:
        wse = Reference(
            ws,
            min_col=8,
            min_row=geom_header_row,
            max_col=7 + len(rows_for),
            max_row=geom_end,
        )
        line.add_data(wse, titles_from_data=True)
    line.set_categories(cats)
    if line.series:
        _style_line_series(line.series[0], "0F172A", width=22000)
        line.series[0].tx = SeriesLabel(v="Ground")
    for offset, scenario in enumerate(rows_for):
        idx = offset + 1
        if idx >= len(line.series):
            break
        _style_line_series(line.series[idx], _hex(scenario.get("color")), width=14000, dash="dash")

    area.y_axis.crosses = "min"
    line.y_axis.axId = 200
    area += line
    ws.add_chart(area, "N5")

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18
    for col in range(3, 16):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.freeze_panes = f"A{geom_start}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.oddHeader.left.text = f"HydroBridge · Transect {transect}"
    return hyd_header_row, geom_end
